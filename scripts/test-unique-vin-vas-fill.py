import sys
import os

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "api"))
from revenue_stream_fill import build_workbook  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
import io  # noqa: E402

payload = {
    "date": "2026-10-24",
    "monthStart": "2026-10-01",
    "monthEnd": "2026-10-24",
    "holidays": [],
    "branches": {
        "CO01A": {"uniqueVinVasForTheDay": 21, "uniqueVinVasAchievementForTheMonth": 617},
    },
}

xlsx_bytes = build_workbook(payload)
wb = load_workbook(io.BytesIO(xlsx_bytes))
ws = wb["Sheet1"]
print("D94 (for the day):", ws["D94"].value)
print("D95 (for the month):", ws["D95"].value)
print("E94 (untouched branch, should stay as template had it or None):", ws["E94"].value)
