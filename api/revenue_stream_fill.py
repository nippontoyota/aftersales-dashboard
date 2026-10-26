"""Vercel Python Function: fills the user's own "Revenue Stream.xlsx"
template (committed at templates/revenue-stream-template.xlsx) with live
dashboard numbers, cell by cell, without touching any formula, style, merge,
or column width already in that file.

Internal-only — called by src/app/api/hq/revenue-stream-export/route.ts
(which does the HQ-admin auth check and gathers the numbers via buildReport(),
report.ts's own aggregation, the same one the live dashboard uses). This
function never talks to the database itself; it only places numbers into
cells (2026-10-23, at the user's request, replacing the from-scratch ExcelJS
version with an exact copy of the template he uses day to day).

Row-by-row mapping (confirmed against the template he supplied,
"Revenue Stream.xlsx", dated 2026-10-20):

- ROW_FIELD_MAP: rows whose branch-column cells are plain literal numbers in
  the template (not formulas) and have a live equivalent on BranchReport.
  Only touched where the template itself already had a literal value for
  that (row, branch) pair — a blank cell there (e.g. the Body & Paint-only
  branches CO01E/TR01B/KL01B have no TGLOSS Service row) means the template
  itself doesn't track that metric for that branch, so it's left blank.
  Exception: ALWAYS_FILL_ROWS (65/66, Scrap/Used Oil Revenue MTD) — no
  branch is actually excluded from scrap/used-oil revenue, so a blank
  template cell there is just a stale gap from 2026-10-20, not an
  exclusion; those two rows get the live value (via fill_row()) regardless
  of what the template cell started as (found 2026-10-26: KT01A, TI01B,
  TI01C and TL01A were silently losing real MTD revenue this way).
- PATCHED_FORMULA_ROWS: rows 73/76 (TGLOSS Service Penetration, TGLOSS SPO)
  hold a VLOOKUP into an external workbook link that doesn't exist alongside
  this file — dead formulas. We already compute the same figures
  (penetrationTGlossService/tGlossSpo), so these two rows get their formula
  replaced with the literal value instead of being left broken.
- UNTRACKED_CLEAR_ROWS: metrics the template tracks (T Club, Extended
  Warranty, Smiles, GR→BP Conversion, TP-TP, NTP-TP) that nothing in our
  upload pipeline captures — same gap the original ExcelJS version
  flagged. Cleared (not left with the template's old demo numbers) so a
  branch isn't misled into thinking a stale Oct-20 figure is today's —
  they're filled in by hand after download, same as before. Unique VIN
  VAS (rows 94/95, 2026-10-24) and Total Car Care Treatments (rows 97/98,
  2026-10-24) both moved out of this list — see unique-vin-vas/compute.ts
  and total-car-care/compute.ts; both are now live ROW_FIELD_MAP entries
  like everything else in the CPU/BPU/Offtake/Parts Retail/PM+OC block.
  Total Car Care Treatments' template labels also had their old
  "(w/o Mech Care)" wording stripped the same day (a direct edit to
  templates/revenue-stream-template.xlsx's shared strings, at the user's
  request) — this metric now counts every TGLOSS treatment, Mech Care
  included.
- CLEAR_FORMULA_ROWS: row 75 (Mech Care Penetration) is the same kind of gap
  as UNTRACKED_CLEAR_ROWS above, but the cell holds the same dead VLOOKUP as
  73/76 rather than a plain number — cleared via patch_formula_row() so the
  broken external-file reference doesn't linger in the formula bar.
- Row 74 (Target TGLOSS Mech Care, a fixed 0.38 constant — the same number as
  row 72's) is left completely untouched: it's not a daily figure, so
  there's nothing to refresh.
- Rows 14-16/18/20/22/24/25/43/46/49/52/93/96/99/102/105-119 and everything
  from row 120 on are either pure formulas (region/group totals, every
  achievement %, every Gentani ratio) or already-blank target placeholders —
  never touched, so Excel recalculates them itself from the inputs above.

Sheet2 ("Service daily report") is a condensed one-row-per-branch view whose
every figure is itself a formula pulling from Sheet1 — but via a broken
external-workbook reference ('[2]Oct 20'!...) rather than an internal one.
fix_sheet2() rewrites those to point at this same workbook's Sheet1 instead,
so the sheet recalculates correctly standalone; its labels, merges, styles
and the handful of VAS-target formulas are left exactly as they are.

Known gap (not fixed here): openpyxl silently drops a small number of this
template's conditional-formatting rules on load (an openpyxl limitation with
some Excel 2010+-only rule formats) — cosmetic highlighting on a dozen-odd
ratio cells may be missing from the output. Flagged to the user; not a data
problem.
"""

from __future__ import annotations

import io
import json
import os
import re
from datetime import date as date_cls
from http.server import BaseHTTPRequestHandler
from typing import Callable, Optional

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "..", "templates", "revenue-stream-template.xlsx")

BRANCH_COLUMNS: dict[str, str] = {
    "CO01A": "D", "CO01B": "E", "CO01E": "F", "MV01A": "G", "KY01A": "H",
    "TR01A": "K", "TR01B": "L", "TR01C": "M", "KL01A": "N", "KL01B": "O", "PH01A": "P",
    "TL01A": "S", "KT01A": "T", "KT01B": "U", "TI01A": "V", "IR01A": "W", "TI01B": "X", "TI01C": "Y",
}
COLUMN_TO_BRANCH: dict[str, str] = {col: branch for branch, col in BRANCH_COLUMNS.items()}

ROW_FIELD_MAP: dict[int, str] = {
    3: "gusRoBilledForTheDay", 4: "gusRoMtd", 5: "gusPartsMtd", 6: "gusLabourMtd",
    8: "bpuRoBilledForTheDay", 9: "bpuRoMtd", 10: "bpuPartsMtd", 11: "bpuLabourMtd",
    13: "externalSalesMtd",
    23: "vasAchievementForTheMonth",
    26: "wheelBalancingForTheDay", 27: "wheelBalancingMtd",
    29: "wheelAlignmentForTheDay", 30: "wheelAlignmentMtd",
    32: "brakeSkimmingForTheDay", 33: "brakeSkimmingMtd",
    41: "engineFlushForTheDay", 42: "engineFlushMtd",
    44: "evaporatorCleaningForTheDay", 45: "evaporatorCleaningMtd",
    47: "diyRevenueForTheDay", 48: "diyRevenueMtd",
    50: "injectorCleanerForTheDay", 51: "injectorCleanerMtd",
    53: "syntheticOilForTheDay", 54: "syntheticOilMtd",
    56: "brakeCleaningSprayForTheDay", 57: "brakeCleaningSprayMtd",
    58: "tireTarget", 59: "tireSales", 60: "tireSalesForTheMonth",
    61: "batteryTarget", 62: "batterySales", 63: "batterySalesForTheMonth",
    65: "scrapRevenueMtd", 66: "usedOilRevenueMtd",
    72: "targetTGlossService",
    79: "cpuForTheDay", 80: "cpuAchievementForTheMonth",
    81: "bpuTarget", 82: "bpuForTheDay", 83: "bpuAchievementForTheMonth",
    84: "offtakeTarget", 85: "offtakeForThePreviousDay", 86: "offtakeAchievementForTheMonth",
    87: "partsRetailTarget", 88: "partsRetailForTheDay", 89: "partsRetailAchievementForTheMonth",
    90: "pmOcTarget", 91: "pmOcForTheDay", 92: "pmOcAchievementForTheMonth",
    94: "uniqueVinVasForTheDay", 95: "uniqueVinVasAchievementForTheMonth",
    97: "totalCarCareTreatmentsForTheDay", 98: "totalCarCareTreatmentsAchievementForTheMonth",
}

# Dead VLOOKUP rows — overwritten even though the original cell is a formula.
# Row 76 is formatted as a percentage (0.0%) and its cached pre-break values
# were all 0-1 ratios, so despite the "T Gloss SPO" label it's the
# achievement ratio (spoTGloss ÷ spoTGlossTarget), not the Rs figure itself.
PATCHED_FORMULA_ROWS: dict[int, str] = {
    73: "penetrationTGlossService",
    76: "tGlossSpo",
}

# Real metrics the template tracks that nothing in our pipeline captures —
# cleared (not left stale) so a branch fills them in by hand. Plain-literal
# rows only; row 75 (Mech Care Penetration) is also untracked but is a dead
# VLOOKUP formula, not a literal, so touch_row() would silently skip it —
# see CLEAR_FORMULA_ROWS below for that one.
UNTRACKED_CLEAR_ROWS = [35, 36, 38, 39, 55, 67, 68, 69, 70, 100, 101, 103, 104]

# Dead VLOOKUP rows with no live equivalent — unlike PATCHED_FORMULA_ROWS,
# these get cleared outright (blank) rather than filled, same reasoning as
# UNTRACKED_CLEAR_ROWS but using patch_formula_row() since the cell holds a
# formula, not a literal. Row 75 ("Penetration Mech Care") pointed at the
# same broken external link as 73/76, but Mech Care Penetration isn't
# anything we track — leaving the old formula in place would keep showing
# whatever date/file it was last pointed at (confirmed 2026-10-23: a user
# opening the generated file saw the Oct-20 external file path in the
# formula bar for this cell).
CLEAR_FORMULA_ROWS = [75]

# Total MTD (Rs) excluding Scrap/Used Oil — matches the template's own row 15
# formula (=Parts+Labour+Parts+Labour+External); rows 17/19 need this same
# sum as a frozen snapshot for a date that isn't otherwise on this sheet.
COMPARISON_TOTAL_FIELDS = ["gusPartsMtd", "gusLabourMtd", "bpuPartsMtd", "bpuLabourMtd", "externalSalesMtd"]

HOLIDAY_CELLS = ["AR69", "AR70", "AR71", "AR72", "AR73"]

EXTERNAL_REF_RE = re.compile(r"'\[\d+\][^']*'!")


def parse_iso_date(value: Optional[str]) -> Optional[date_cls]:
    if not value:
        return None
    y, m, d = value.split("-")
    return date_cls(int(y), int(m), int(d))


def is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith("=")


def touch_row(ws, row: int, value_for_branch: Callable[[str], Optional[float]]) -> None:
    """Overwrites a branch's cell only if the template already had a plain
    (non-formula) value there — skips formulas and skips branches the
    template's own layout leaves blank for this row."""
    for branch, col in BRANCH_COLUMNS.items():
        cell = ws[f"{col}{row}"]
        if cell.value is None or is_formula(cell.value):
            continue
        cell.value = value_for_branch(branch)


# Rows 65/66 (Scrap Revenue MTD, Used Oil Revenue MTD) apply to every branch —
# unlike the Body & Paint-only exclusions elsewhere on this sheet, no branch
# is structurally excluded from scrap/used-oil revenue. A blank cell here is
# just whatever the template happened to have on 2026-10-20 (some branches
# hadn't billed scrap/oil yet that day), not a "this branch doesn't track
# this" signal — so these two rows get the live value regardless of what the
# template cell started as (2026-10-26, fixed after KT01A/TI01B/TI01C/TL01A
# were found missing real revenue because their template cells were blank).
ALWAYS_FILL_ROWS = {65, 66}


def fill_row(ws, row: int, value_for_branch: Callable[[str], Optional[float]]) -> None:
    """Like touch_row, but writes every branch's cell regardless of whether
    the template cell started blank — only an existing formula is left
    alone."""
    for branch, col in BRANCH_COLUMNS.items():
        cell = ws[f"{col}{row}"]
        if is_formula(cell.value):
            continue
        cell.value = value_for_branch(branch)


# Rows 73/75/76's dead VLOOKUP was dragged across the whole D:AB range when
# the template was built, not just the 18 branch columns — it's also sitting
# in the region-total (I/Q/Z), group-total (AB) and spacer (J/R/AA) columns.
DEAD_VLOOKUP_COLUMN_RANGE = range(4, 29)  # D .. AB


def patch_formula_row(ws, row: int, value_for_branch: Callable[[str], Optional[float]]) -> None:
    """Like touch_row, but also overwrites formula cells, and covers every
    column D:AB rather than just the branch columns — for the dead VLOOKUP
    rows (73/75/76), which is a wider range than every other mapped row.
    A non-branch column (region/group totals, spacers) always clears to
    None: these are ratio metrics, which the live dashboard itself leaves
    blank at the total level (see daily-report-rows.ts's regionTotalCell)."""
    for col_idx in DEAD_VLOOKUP_COLUMN_RANGE:
        col = get_column_letter(col_idx)
        cell = ws[f"{col}{row}"]
        if cell.value is None:
            continue
        branch = COLUMN_TO_BRANCH.get(col)
        cell.value = value_for_branch(branch) if branch is not None else None


def comparison_total(branch: str, branches: Optional[dict]) -> Optional[float]:
    if not branches:
        return None
    bd = branches.get(branch)
    if not bd:
        return None
    parts = [bd.get(f) for f in COMPARISON_TOTAL_FIELDS]
    if any(p is None for p in parts):
        return None
    return sum(parts)


def fix_sheet2(ws2) -> None:
    """Rewrites Sheet2's broken external-workbook references ('[2]Oct 20'!…)
    to point at this same workbook's Sheet1 instead, so it recalculates
    standalone. Everything else on Sheet2 (labels, merges, styles, the
    handful of VAS-target formulas) is left untouched."""
    for row in ws2.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("=") and EXTERNAL_REF_RE.search(v):
                cell.value = EXTERNAL_REF_RE.sub("Sheet1!", v)


def build_workbook(payload: dict) -> bytes:
    wb = load_workbook(TEMPLATE_PATH, data_only=False)
    ws1 = wb["Sheet1"]
    ws2 = wb["Sheet2"]

    branches_today: dict = payload.get("branches") or {}
    branches_last_month: Optional[dict] = payload.get("branchesLastMonth")
    branches_last_year: Optional[dict] = payload.get("branchesLastYear")

    for row, field in ROW_FIELD_MAP.items():
        writer = fill_row if row in ALWAYS_FILL_ROWS else touch_row
        writer(ws1, row, lambda b, f=field: (branches_today.get(b) or {}).get(f))

    for row, field in PATCHED_FORMULA_ROWS.items():
        patch_formula_row(ws1, row, lambda b, f=field: (branches_today.get(b) or {}).get(f))

    for row in UNTRACKED_CLEAR_ROWS:
        touch_row(ws1, row, lambda b: None)

    for row in CLEAR_FORMULA_ROWS:
        patch_formula_row(ws1, row, lambda b: None)

    touch_row(ws1, 17, lambda b: comparison_total(b, branches_last_month))
    touch_row(ws1, 19, lambda b: comparison_total(b, branches_last_year))

    ws1["A2"] = parse_iso_date(payload.get("date"))
    if payload.get("lastMonthDate"):
        ws1["A17"] = parse_iso_date(payload["lastMonthDate"])
    if payload.get("lastYearDate"):
        ws1["A19"] = parse_iso_date(payload["lastYearDate"])

    ws1["AP69"] = parse_iso_date(payload.get("monthStart"))
    ws1["AQ69"] = parse_iso_date(payload.get("monthEnd"))
    holidays = payload.get("holidays") or []
    for i, cell_ref in enumerate(HOLIDAY_CELLS):
        ws1[cell_ref] = parse_iso_date(holidays[i]) if i < len(holidays) else None

    fix_sheet2(ws2)

    # The template itself opens to Sheet2 by default (its workbook-level
    # active-tab setting, independent of sheet order) — overridden to Sheet1
    # per the user's request (2026-10-23).
    wb.active = wb.sheetnames.index("Sheet1")
    for i, ws in enumerate(wb.worksheets):
        ws.sheet_view.tabSelected = i == wb.sheetnames.index("Sheet1")

    wb.calculation.fullCalcOnLoad = True

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class handler(BaseHTTPRequestHandler):
    def do_POST(self):
        try:
            length = int(self.headers.get("content-length", 0))
            raw = self.rfile.read(length)
            payload = json.loads(raw)
            xlsx_bytes = build_workbook(payload)
        except Exception as exc:  # noqa: BLE001 — surfaced to the caller as JSON, not swallowed
            body = json.dumps({"error": str(exc)}).encode("utf-8")
            self.send_response(500)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return

        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Length", str(len(xlsx_bytes)))
        self.end_headers()
        self.wfile.write(xlsx_bytes)
